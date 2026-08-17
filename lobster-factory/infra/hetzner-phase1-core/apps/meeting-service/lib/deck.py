"""minutes dict -> PPTX。

品牌色複用 next-admin lib/theme/tokens.ts 的預設 token。
"""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ACCENT = RGBColor(0x63, 0x66, 0xF1)
INK = RGBColor(0x0D, 0x15, 0x26)
MUTED = RGBColor(0x47, 0x55, 0x69)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
BODY_BG = RGBColor(0xF2, 0xF5, 0xFB)
BORDER = RGBColor(0xE6, 0xEA, 0xF4)
GREEN = RGBColor(0x10, 0xB9, 0x81)
RED = RGBColor(0xEF, 0x44, 0x44)
YELLOW = RGBColor(0xF5, 0x9E, 0x0B)

FONT_EA = "Microsoft JhengHei"
FONT_LAT = "Segoe UI"
PRIORITY_COLOR = {"high": RED, "medium": YELLOW, "low": GREEN}


def set_font(run, size=18, bold=False, color=INK):
    """python-pptx 只設 latin typeface；CJK 要另外寫 <a:ea> 否則會掉回新細明體。"""
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    f.name = FONT_LAT
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", FONT_EA)


def blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = BODY_BG
    return s


def add_text(slide, left, top, width, height, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].alignment = align
    return tf


def accent_bar(slide, top=0.62):
    bar = slide.shapes.add_shape(1, Inches(0.62), Inches(top), Inches(0.9), Inches(0.055))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    bar.shadow.inherit = False


def header(slide, title, kicker=None):
    if kicker:
        tf = add_text(slide, 0.62, 0.28, 11, 0.32)
        set_font(tf.paragraphs[0].add_run(), 12, True, ACCENT)
        tf.paragraphs[0].runs[0].text = kicker.upper()
    accent_bar(slide)
    tf = add_text(slide, 0.62, 0.78, 11.5, 0.9)
    r = tf.paragraphs[0].add_run()
    r.text = title
    set_font(r, 30, True, INK)


def card(slide, left, top, width, height, fill=CARD):
    sh = slide.shapes.add_shape(1, Inches(left), Inches(top), Inches(width), Inches(height))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = BORDER
    sh.line.width = Pt(1)
    sh.shadow.inherit = False
    return sh


def bullets(slide, items, left=0.62, top=1.9, width=11.5, size=17, gap=0.52):
    for i, it in enumerate(items):
        tf = add_text(slide, left, top + i * gap, width, gap)
        p = tf.paragraphs[0]
        d = p.add_run()
        d.text = "▪  "
        set_font(d, size, True, ACCENT)
        r = p.add_run()
        r.text = str(it)
        set_font(r, size, False, INK)


def slide_title(prs, m):
    s = blank(prs)
    band = s.shapes.add_shape(1, Inches(0), Inches(0), prs.slide_width, Inches(3.05))
    band.fill.solid()
    band.fill.fore_color.rgb = ACCENT
    band.line.fill.background()
    band.shadow.inherit = False

    tf = add_text(s, 0.9, 0.95, 11.5, 1.2)
    r = tf.paragraphs[0].add_run()
    r.text = m.get("title", "會議記錄")
    set_font(r, 40, True, CARD)

    meta = "　|　".join(x for x in [
        m.get("date", ""),
        f"{m.get('duration_min', '?')} 分鐘",
        f"{len(m.get('attendees', []))} 位與會者",
    ] if x)
    tf = add_text(s, 0.92, 2.15, 11.5, 0.5)
    r = tf.paragraphs[0].add_run()
    r.text = meta
    set_font(r, 15, False, RGBColor(0xE0, 0xE7, 0xFF))

    if m.get("attendees"):
        tf = add_text(s, 0.9, 3.5, 11.5, 0.4)
        r = tf.paragraphs[0].add_run()
        r.text = "與會：" + "、".join(m["attendees"])
        set_font(r, 14, False, MUTED)
    return s


def slide_tldr(prs, m):
    s = blank(prs)
    header(s, "三句話結論", "TL;DR")
    for i, t in enumerate(m.get("tldr", [])[:3]):
        card(s, 0.62, 2.0 + i * 1.35, 11.9, 1.12)
        tf = add_text(s, 0.95, 2.2 + i * 1.35, 0.6, 0.7)
        r = tf.paragraphs[0].add_run()
        r.text = str(i + 1)
        set_font(r, 30, True, ACCENT)
        tf = add_text(s, 1.65, 2.28 + i * 1.35, 10.5, 0.8)
        r = tf.paragraphs[0].add_run()
        r.text = str(t)
        set_font(r, 17, False, INK)
    return s


def slide_decisions(prs, m):
    d = m.get("decisions", [])
    if not d:
        return None
    s = blank(prs)
    header(s, "決議事項", f"{len(d)} DECISIONS")
    top = 1.95
    for it in d[:5]:
        card(s, 0.62, top, 11.9, 1.0)
        chip = s.shapes.add_shape(1, Inches(0.85), Inches(top + 0.18), Inches(0.16), Inches(0.64))
        chip.fill.solid()
        chip.fill.fore_color.rgb = GREEN
        chip.line.fill.background()
        chip.shadow.inherit = False
        tf = add_text(s, 1.2, top + 0.14, 8.6, 0.4)
        r = tf.paragraphs[0].add_run()
        r.text = it.get("what", "")
        set_font(r, 16, True, INK)
        tf = add_text(s, 1.2, top + 0.55, 8.6, 0.35)
        r = tf.paragraphs[0].add_run()
        r.text = "理由：" + it.get("rationale", "—")
        set_font(r, 12, False, MUTED)
        tf = add_text(s, 10.0, top + 0.32, 2.3, 0.4)
        r = tf.paragraphs[0].add_run()
        r.text = it.get("owner", "")
        set_font(r, 13, True, ACCENT)
        top += 1.12
    return s


def slide_actions(prs, m):
    a = m.get("actions", [])
    if not a:
        return None
    s = blank(prs)
    header(s, "待辦與負責人", f"{len(a)} ACTION ITEMS")
    rows = min(len(a), 7) + 1
    tbl = s.shapes.add_table(rows, 4, Inches(0.62), Inches(1.95),
                             Inches(11.9), Inches(0.42 * rows)).table
    for w, c in zip((5.9, 2.0, 2.0, 2.0), range(4)):
        tbl.columns[c].width = Inches(w)
    for c, h in enumerate(("待辦事項", "負責人", "期限", "優先")):
        cell = tbl.cell(0, c)
        cell.text = ""
        cell.fill.solid()
        cell.fill.fore_color.rgb = ACCENT
        r = cell.text_frame.paragraphs[0].add_run()
        r.text = h
        set_font(r, 13, True, CARD)
    for i, it in enumerate(a[:7], 1):
        pri = str(it.get("priority", "medium")).lower()
        vals = (it.get("task", ""), it.get("owner", "—"),
                it.get("due", "—"), pri.upper())
        for c, v in enumerate(vals):
            cell = tbl.cell(i, c)
            cell.text = ""
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD
            p = cell.text_frame.paragraphs[0]
            r = p.add_run()
            r.text = str(v)
            set_font(r, 12, c == 3, PRIORITY_COLOR.get(pri, MUTED) if c == 3 else INK)
    return s


def slide_topic(prs, t):
    s = blank(prs)
    header(s, t.get("heading", ""), "討論主題")
    bullets(s, t.get("points", [])[:5])
    q = (t.get("quotes") or [None])[0]
    if q:
        card(s, 0.62, 5.1, 11.9, 1.15, RGBColor(0xE0, 0xF2, 0xFE))
        tf = add_text(s, 1.0, 5.32, 11.1, 0.8)
        r = tf.paragraphs[0].add_run()
        r.text = f"「{q.get('text', '')}」"
        set_font(r, 14, False, RGBColor(0x03, 0x69, 0xA1))
        r2 = tf.paragraphs[0].add_run()
        r2.text = f"   — {q.get('speaker', '')} {q.get('t', '')}"
        set_font(r2, 11, False, MUTED)
    return s


def slide_open(prs, m):
    items = m.get("open_questions", [])
    risks = m.get("risks", [])
    if not items and not risks:
        return None
    s = blank(prs)
    header(s, "未決事項與風險", "OPEN")
    if items:
        card(s, 0.62, 1.95, 5.85, 4.2)
        tf = add_text(s, 0.95, 2.15, 5.2, 0.4)
        r = tf.paragraphs[0].add_run()
        r.text = "❓ 尚待確認"
        set_font(r, 15, True, ACCENT)
        bullets(s, items[:5], left=0.95, top=2.7, width=5.2, size=13, gap=0.62)
    if risks:
        card(s, 6.67, 1.95, 5.85, 4.2)
        tf = add_text(s, 7.0, 2.15, 5.2, 0.4)
        r = tf.paragraphs[0].add_run()
        r.text = "⚠ 風險"
        set_font(r, 15, True, RED)
        bullets(s, risks[:5], left=7.0, top=2.7, width=5.2, size=13, gap=0.62)
    return s


def plan_slides(minutes: dict) -> list[str]:
    """算出這份 minutes 會產生哪些頁面（順序固定），供講稿生成對齊頁數用。
    build() 與這裡必須同步——加減頁面時兩處一起改。"""
    plan = ["title", "tldr"]
    if minutes.get("decisions"):
        plan.append("decisions")
    if minutes.get("actions"):
        plan.append("actions")
    for t in minutes.get("topics", [])[:6]:
        plan.append(f"topic:{t.get('heading','')}")
    if minutes.get("open_questions") or minutes.get("risks"):
        plan.append("open")
    return plan


def build(minutes: dict, out_path: str) -> int:
    """回傳 slide 數。順序須與 plan_slides() 一致。"""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide_title(prs, minutes)
    slide_tldr(prs, minutes)
    if minutes.get("decisions"):
        slide_decisions(prs, minutes)
    if minutes.get("actions"):
        slide_actions(prs, minutes)
    for t in minutes.get("topics", [])[:6]:
        slide_topic(prs, t)
    if minutes.get("open_questions") or minutes.get("risks"):
        slide_open(prs, minutes)

    prs.save(out_path)
    return len(prs.slides._sldIdLst)
