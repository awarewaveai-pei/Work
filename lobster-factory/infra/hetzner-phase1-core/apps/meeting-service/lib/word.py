"""minutes dict -> .docx 會議記錄。

為什麼有這支（跟 minutes_md.py 的分工）：`.md` 是給存檔／貼進 repo 或 wiki 用的，
但客戶端傳閱會議記錄的實際格式是 Word——對方不會裝 markdown 檢視器，也常需要
直接在檔案上加註、追蹤修訂、轉 PDF 寄出。內容結構刻意與 minutes_md.render()
對齊（同樣的六個章節 + 附錄），差別只在呈現載體，避免兩份產出講不同的話。

品牌色與字型沿用 deck.py 的 token（同一份 next-admin lib/theme/tokens.ts 預設），
不另立一套。CJK 字型要手動寫 <w:rFonts w:eastAsia>，理由同 deck.py 的 set_font()：
python-docx 預設只設 ascii/hAnsi，中文會掉回新細明體。
"""
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ACCENT = RGBColor(0x63, 0x66, 0xF1)
INK = RGBColor(0x0D, 0x15, 0x26)
MUTED = RGBColor(0x47, 0x55, 0x69)
RED = RGBColor(0xEF, 0x44, 0x44)
YELLOW = RGBColor(0xB4, 0x7A, 0x08)  # md/pptx 的黃在白底 Word 上對比不足，壓深
GREEN = RGBColor(0x0F, 0x8A, 0x62)

FONT_EA = "Microsoft JhengHei"
FONT_LAT = "Segoe UI"

PRI = {"high": ("高", RED), "medium": ("中", YELLOW), "low": ("低", GREEN)}


def _style_run(run, size=11, bold=False, color=INK, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = FONT_LAT
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), FONT_EA)
    rFonts.set(qn("w:ascii"), FONT_LAT)
    rFonts.set(qn("w:hAnsi"), FONT_LAT)


def _para(doc, text="", size=11, bold=False, color=INK, italic=False,
          space_after=6, align=None, style=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        _style_run(p.add_run(text), size=size, bold=bold, color=color, italic=italic)
    return p


def _heading(doc, text, size=15):
    p = _para(doc, text, size=size, bold=True, color=ACCENT, space_after=4)
    p.paragraph_format.space_before = Pt(14)
    return p


def _table(doc, headers, rows):
    """rows 的每格可以是 str，或 (str, RGBColor) 指定該格文字顏色。"""
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        _style_run(cell.paragraphs[0].add_run(h), size=10, bold=True, color=ACCENT)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            text, color = val if isinstance(val, tuple) else (val, INK)
            cells[i].text = ""
            _style_run(cells[i].paragraphs[0].add_run(str(text)), size=10, color=color)
    return t


def build(m: dict, out_path: str) -> str:
    doc = Document()
    # 預設樣式也要設 eastAsia，否則沒被 _style_run 走到的地方仍會掉字型。
    # 用 get_or_add_* 而非直接 .rPr.rFonts：後者只有在前面剛好設過 font.name
    # （順帶建出 rPr/rFonts 節點）時才不會是 None，等於隱性依賴敘述順序，
    # 有人調換兩行就 AttributeError。與 _style_run 的作法保持一致。
    normal = doc.styles["Normal"]
    normal.font.name = FONT_LAT
    normal.font.size = Pt(11)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT_EA)

    _para(doc, m.get("title", "會議記錄"), size=22, bold=True, color=INK,
          space_after=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    dec = m.get("decisions", []) or []
    act = m.get("actions", []) or []
    _table(doc, ["項目", "內容"], [
        ["日期", m.get("date") or "—"],
        ["時長", f"{m.get('duration_min','—')} 分鐘"],
        ["出席", "、".join(m.get("attendees", []) or []) or "—"],
        ["產出", f"{len(dec)} 項決議、{len(act)} 項待辦"],
    ])

    if m.get("_provenance"):
        _para(doc, "")
        p = _para(doc, "資料來源與限制：", size=9, bold=True, color=MUTED, space_after=2)
        _style_run(p.add_run(m["_provenance"]), size=9, color=MUTED, italic=True)

    _heading(doc, "一、三句話結論")
    tldr = m.get("tldr", []) or []
    if tldr:
        for t in tldr:
            _para(doc, t, style="List Number", space_after=3)
    else:
        _para(doc, "（本次錄音內容不構成正式會議結論）", color=MUTED, italic=True)

    if dec:
        _heading(doc, "二、決議事項")
        _table(doc, ["#", "決議", "負責", "理由"],
               [[i, d.get("what", ""), d.get("owner", "—"), d.get("rationale", "—")]
                for i, d in enumerate(dec, 1)])

    if act:
        _heading(doc, "三、待辦事項")
        rows = []
        for i, a in enumerate(act, 1):
            label, color = PRI.get(str(a.get("priority", "medium")).lower(),
                                   (a.get("priority", "—"), INK))
            rows.append([i, a.get("task", ""), a.get("owner", "—"),
                         a.get("due", "—"), (label, color)])
        _table(doc, ["#", "待辦", "負責人", "期限", "優先"], rows)

    if m.get("topics"):
        _heading(doc, "四、討論內容")
        for i, t in enumerate(m["topics"], 1):
            _para(doc, f"4.{i}　{t.get('heading','')}", size=12, bold=True, space_after=3)
            for pt in t.get("points", []) or []:
                _para(doc, pt, style="List Bullet", space_after=2)
            for q in t.get("quotes", []) or []:
                qp = _para(doc, f"「{q.get('text','')}」", size=10, color=MUTED,
                           italic=True, space_after=1)
                qp.paragraph_format.left_indent = Pt(24)
                sp = _para(doc, f"— {q.get('speaker','')}　{q.get('t','')}",
                           size=9, color=MUTED, space_after=6)
                sp.paragraph_format.left_indent = Pt(24)

    if m.get("open_questions"):
        _heading(doc, "五、未決事項")
        for q in m["open_questions"]:
            _para(doc, f"☐ {q}", style="List Bullet", space_after=2)

    if m.get("risks"):
        _heading(doc, "六、風險")
        for r in m["risks"]:
            _para(doc, r, style="List Number", color=RED, space_after=2)

    c = m.get("coaching", {}) or {}
    per = c.get("per_slide", []) or []
    if per:
        doc.add_page_break()
        _heading(doc, "附錄：簡報大綱")
        total = sum(x.get("seconds", 0) for x in per)
        _para(doc, f"共 {len(per)} 頁，目標時長 {total//60} 分 {total%60} 秒。"
                   f"完整講稿見報告教練檔與 PPT 備忘稿。", size=10, color=MUTED)
        _table(doc, ["頁", "這頁要達成什麼", "秒數"],
               [[i, n.get("goal", ""), n.get("seconds", 0)] for i, n in enumerate(per, 1)])

    doc.save(out_path)
    return out_path
