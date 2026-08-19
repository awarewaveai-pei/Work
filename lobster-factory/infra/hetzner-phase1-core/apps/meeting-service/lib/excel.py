"""minutes dict -> 固定規格 Excel（.xlsx），跟 PPT/md 同層級的額外輸出。

五個工作表，欄位對應 minutes.json 的完整結構：
  摘要      — title/date/duration/attendees/tldr/_provenance
  決議事項  — decisions[].{what,owner,rationale}
  待辦事項  — actions[].{task,owner,due,priority}（含優先級色塊）
  討論主題  — topics[].{heading,points,quotes}
  未決與風險 — open_questions / risks

品牌色複用 lib/deck.py 的同一組 token，跟 PPT 視覺一致。
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

ACCENT = "6366F1"
INK = "0D1526"
MUTED = "475569"
BORDER_C = "E6EAF4"
GREEN = "10B981"
RED = "EF4444"
YELLOW = "F59E0B"
BODY_BG = "F2F5FB"

HEADER_FONT = Font(name="Microsoft JhengHei", size=11, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=ACCENT)
TITLE_FONT = Font(name="Microsoft JhengHei", size=16, bold=True, color=INK)
LABEL_FONT = Font(name="Microsoft JhengHei", size=10, bold=True, color=MUTED)
BODY_FONT = Font(name="Microsoft JhengHei", size=10, color=INK)
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Side(style="thin", color=BORDER_C)
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

PRIORITY_FILL = {
    "high": PatternFill("solid", fgColor=RED),
    "medium": PatternFill("solid", fgColor=YELLOW),
    "low": PatternFill("solid", fgColor=GREEN),
}


def _header_row(ws: Worksheet, row: int, headers: list, widths: list = None):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = CELL_BORDER
    if widths:
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = w


def _data_row(ws: Worksheet, row: int, values: list, wrap_cols: set = None):
    wrap_cols = wrap_cols or set()
    for c, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=c, value=v if v not in (None, "") else "—")
        cell.font = BODY_FONT
        cell.border = CELL_BORDER
        if c in wrap_cols:
            cell.alignment = WRAP
        else:
            cell.alignment = Alignment(vertical="top")


def _sheet_summary(wb: Workbook, m: dict):
    ws = wb.active
    ws.title = "摘要"
    ws.sheet_view.showGridLines = False

    ws["A1"] = m.get("title", "會議記錄")
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:D1")
    ws.row_dimensions[1].height = 28

    meta = [
        ("日期", m.get("date") or "—"),
        ("時長", f"{m.get('duration_min', '—')} 分鐘"),
        ("出席", "、".join(m.get("attendees", [])) or "—"),
        ("決議數", str(len(m.get("decisions", [])))),
        ("待辦數", str(len(m.get("actions", [])))),
    ]
    r = 3
    for label, val in meta:
        ws.cell(row=r, column=1, value=label).font = LABEL_FONT
        c = ws.cell(row=r, column=2, value=val)
        c.font = BODY_FONT
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="三句話結論").font = LABEL_FONT
    r += 1
    for t in m.get("tldr", []):
        ws.cell(row=r, column=1, value=f"• {t}").font = BODY_FONT
        ws.cell(row=r, column=1).alignment = WRAP
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        ws.row_dimensions[r].height = 30
        r += 1
    if not m.get("tldr"):
        ws.cell(row=r, column=1, value="（無——此錄音判定非正式會議或內容不足）").font = BODY_FONT
        r += 1

    if m.get("_provenance"):
        r += 1
        ws.cell(row=r, column=1, value="資料來源與限制").font = LABEL_FONT
        r += 1
        ws.cell(row=r, column=1, value=m["_provenance"]).font = Font(
            name="Microsoft JhengHei", size=9, italic=True, color=MUTED)
        ws.cell(row=r, column=1).alignment = WRAP
        ws.merge_cells(start_row=r, start_column=1, end_row=r + 3, end_column=4)
        ws.row_dimensions[r].height = 20

    for col, w in zip("ABCD", (14, 30, 20, 20)):
        ws.column_dimensions[col].width = w


def _sheet_decisions(wb: Workbook, m: dict):
    ws = wb.create_sheet("決議事項")
    ws.sheet_view.showGridLines = False
    _header_row(ws, 1, ["#", "決議", "負責人", "理由"], widths=[5, 40, 15, 45])
    decisions = m.get("decisions", [])
    if not decisions:
        ws.cell(row=2, column=1, value="（本次無決議事項）").font = BODY_FONT
        ws.merge_cells("A2:D2")
        return
    for i, d in enumerate(decisions, 1):
        _data_row(ws, i + 1, [i, d.get("what"), d.get("owner"), d.get("rationale")],
                  wrap_cols={2, 4})
        ws.row_dimensions[i + 1].height = 45


def _sheet_actions(wb: Workbook, m: dict):
    ws = wb.create_sheet("待辦事項")
    ws.sheet_view.showGridLines = False
    _header_row(ws, 1, ["#", "待辦事項", "負責人", "期限", "優先級"], widths=[5, 45, 15, 15, 12])
    actions = m.get("actions", [])
    if not actions:
        ws.cell(row=2, column=1, value="（本次無待辦事項）").font = BODY_FONT
        ws.merge_cells("A2:E2")
        return
    for i, a in enumerate(actions, 1):
        row = i + 1
        pri = str(a.get("priority", "medium")).lower()
        _data_row(ws, row, [i, a.get("task"), a.get("owner"), a.get("due"), pri.upper()],
                  wrap_cols={2})
        pf = PRIORITY_FILL.get(pri)
        if pf:
            cell = ws.cell(row=row, column=5)
            cell.fill = pf
            cell.font = Font(name="Microsoft JhengHei", size=10, bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[row].height = 32


def _sheet_topics(wb: Workbook, m: dict):
    ws = wb.create_sheet("討論主題")
    ws.sheet_view.showGridLines = False
    _header_row(ws, 1, ["#", "主題", "重點", "引言", "講者", "時間"], widths=[5, 25, 45, 40, 10, 10])
    topics = m.get("topics", [])
    if not topics:
        ws.cell(row=2, column=1, value="（本次無討論主題記錄）").font = BODY_FONT
        ws.merge_cells("A2:F2")
        return
    r = 2
    for i, t in enumerate(topics, 1):
        points = "\n".join(f"• {p}" for p in t.get("points", []))
        quotes = t.get("quotes") or [{}]
        first = quotes[0] if quotes else {}
        _data_row(ws, r, [i, t.get("heading"), points, first.get("text"),
                          first.get("speaker"), first.get("t")], wrap_cols={2, 3, 4})
        ws.row_dimensions[r].height = max(20 * len(t.get("points", [])), 40)
        r += 1
        for extra in quotes[1:]:
            _data_row(ws, r, ["", "", "", extra.get("text"), extra.get("speaker"), extra.get("t")],
                      wrap_cols={4})
            r += 1


def _sheet_open(wb: Workbook, m: dict):
    ws = wb.create_sheet("未決與風險")
    ws.sheet_view.showGridLines = False
    _header_row(ws, 1, ["類型", "內容"], widths=[15, 70])
    r = 2
    for q in m.get("open_questions", []):
        _data_row(ws, r, ["❓ 未決事項", q], wrap_cols={2})
        ws.row_dimensions[r].height = 28
        r += 1
    for risk in m.get("risks", []):
        _data_row(ws, r, ["⚠ 風險", risk], wrap_cols={2})
        cell = ws.cell(row=r, column=1)
        cell.font = Font(name="Microsoft JhengHei", size=10, bold=True, color=RED)
        ws.row_dimensions[r].height = 28
        r += 1
    if r == 2:
        ws.cell(row=2, column=1, value="（無未決事項或風險記錄）").font = BODY_FONT
        ws.merge_cells("A2:B2")


def build(minutes: dict, out_path: str) -> str:
    wb = Workbook()
    _sheet_summary(wb, minutes)
    _sheet_decisions(wb, minutes)
    _sheet_actions(wb, minutes)
    _sheet_topics(wb, minutes)
    _sheet_open(wb, minutes)
    wb.save(out_path)
    return out_path
