"""把講稿寫進 PPTX 備忘稿欄位。

原本這支還負責輸出報告教練 markdown（`write_coach_md()`），2026-08-27 操作者
決定文件類輸出一律用 Word，該函式已移除，改由 `word.build_coach()` 產出
`coach.docx`。刪除而非保留是刻意的——兩套渲染同一份資料必然漂移。
"""
from pptx import Presentation

WPM_ZH = 240  # 中文口語每分鐘字數，正式簡報建議 220-260（word.py 有同值副本）


def attach_notes(pptx_path: str, coaching: dict, out_path: str = None):
    prs = Presentation(pptx_path)
    slides = list(prs.slides)
    notes = coaching.get("per_slide", [])
    # 講稿以位置索引對應投影片。頁數一變就整體錯位，而且是靜默錯位
    # （第 N 頁拿到第 M 頁的稿）。數量不符一律硬失敗，不允許默默對錯。
    if len(notes) != len(slides):
        raise ValueError(
            f"講稿 {len(notes)} 段 != 投影片 {len(slides)} 頁；"
            f"位置索引會靜默錯位，請補齊後重跑"
        )
    for i, s in enumerate(slides):
        n = notes[i]
        body = []
        if n.get("goal"):
            body.append(f"【這頁要達成什麼】{n['goal']}")
        if n.get("script"):
            body.append(f"\n【講稿】\n{n['script']}")
        if n.get("seconds"):
            chars = len(n.get("script", ""))
            est = chars / WPM_ZH * 60
            body.append(f"\n【時間】目標 {n['seconds']} 秒（講稿 {chars} 字，語速 {WPM_ZH} 字/分 約需 {est:.0f} 秒）")
        if n.get("transition"):
            body.append(f"\n【轉場】{n['transition']}")
        if n.get("watch_out"):
            body.append(f"\n【小心】{n['watch_out']}")
        s.notes_slide.notes_text_frame.text = "\n".join(body)
    out = out_path or pptx_path
    prs.save(out)
    return out


